"""TSV/CSV parse+serialize and plate-header stripping — mirrors the `TSV` and `CSV` enums
in `Sources/PLayout/IO/TableIO.swift`. The tidy grid and workbook builder (`Exporter`)
join this module in M4. Pure Python."""
from __future__ import annotations

from playout.model.plate_format import WellNaming

Grid = list[list[str]]


def _pad(rows: list[list[str]]) -> Grid:
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


class TSV:
    """Tab-separated text — what Excel puts on the pasteboard, so the interchange
    format for copy/paste in both directions."""

    @staticmethod
    def serialize(grid: Grid) -> str:
        """Cells joined by tabs, rows by newlines, **no trailing newline**, no quoting."""
        return "\n".join("\t".join(row) for row in grid)

    @staticmethod
    def parse(text: str) -> Grid:
        normalized = text.replace("\r\n", "\n").replace("\r", "\n")
        lines = normalized.split("\n")
        while lines and lines[-1] == "":
            lines.pop()
        if not lines:
            return []
        return _pad([line.split("\t") for line in lines])

    @staticmethod
    def stripping_plate_headers(grid: Grid) -> Grid:
        """Excel users routinely copy a plate map *including* its A–H / 1–12 headers.
        Strip them only when the grid genuinely looks like one: blank corner, every other
        top-row cell blank or its 1-based column number, every later first cell blank or
        the row letter of its 0-based offset."""
        if len(grid) < 2 or len(grid[0]) < 2:
            return grid
        first = grid[0]
        corner_empty = first[0].strip(" \t") == ""

        def looks_numeric(offset: int, cell: str) -> bool:
            t = cell.strip(" \t")
            if t == "":
                return True
            try:
                return int(t) == offset + 1
            except ValueError:
                return False

        top_numeric = all(looks_numeric(i, cell) for i, cell in enumerate(first[1:]))
        left_alpha = all(
            (row[0].strip(" \t") if row else "") == "" or WellNaming.row_index(row[0].strip(" \t")) == offset
            for offset, row in enumerate(grid[1:])
        )
        if not (corner_empty and top_numeric and left_alpha):
            return grid
        return [row[1:] for row in grid[1:]]


class CSV:
    @staticmethod
    def serialize(grid: Grid) -> str:
        """RFC-4180-ish: quote only when needed, double inner quotes, `\\n` line ends,
        **with a trailing newline**."""
        return "\n".join(",".join(CSV._field(cell) for cell in row) for row in grid) + "\n"

    @staticmethod
    def _field(s: str) -> str:
        if not any(ch in s for ch in ',"\n\r'):
            return s
        return '"' + s.replace('"', '""') + '"'

    @staticmethod
    def parse(text: str) -> Grid:
        rows: list[list[str]] = []
        row: list[str] = []
        field = ""
        in_quotes = False
        i = 0
        n = len(text)
        while i < n:
            ch = text[i]
            i += 1
            if in_quotes:
                if ch == '"':
                    if i < n:
                        peek = text[i]
                        i += 1
                        if peek == '"':
                            field += '"'
                        else:
                            in_quotes = False
                            i -= 1  # re-read `peek` outside quotes
                    else:
                        in_quotes = False
                else:
                    field += ch
            else:
                if ch == '"':
                    in_quotes = True
                elif ch == ",":
                    row.append(field)
                    field = ""
                elif ch == "\n":
                    row.append(field)
                    field = ""
                    rows.append(row)
                    row = []
                elif ch == "\r":
                    continue
                else:
                    field += ch
        if field or row:
            row.append(field)
            rows.append(row)
        while rows and all(cell == "" for cell in rows[-1]):
            rows.pop()
        return _pad(rows)


# ======================================================================================
# Workbook / tidy-table construction — mirrors `Exporter` in TableIO.swift and the cell
# styling of XLSXWriter.swift, written through openpyxl (PORT.md §A6).

from enum import Enum  # noqa: E402

from playout.model import palette  # noqa: E402
from playout.model.layout import Factor, FactorKind, Layout, Plate  # noqa: E402


class WorkbookLayout(str, Enum):
    sheetPerFactor = "sheetPerFactor"
    allFactorsOneSheet = "allFactorsOneSheet"

    @property
    def label(self) -> str:
        return "One sheet per factor" if self is WorkbookLayout.sheetPerFactor else "All factors on one sheet"

    @property
    def detail(self) -> str:
        if self is WorkbookLayout.sheetPerFactor:
            return "A separate tab for each factor, ready to paste elsewhere."
        return "One tab per plate, each map headed by its factor name."

    @classmethod
    def lenient(cls, raw) -> "WorkbookLayout":
        try:
            return cls(raw)
        except ValueError:
            return cls.sheetPerFactor


class WorkbookScope(str, Enum):
    allPlates = "allPlates"
    activePlate = "activePlate"

    @classmethod
    def lenient(cls, raw) -> "WorkbookScope":
        try:
            return cls(raw)
        except ValueError:
            return cls.allPlates


JOINT_FALLBACK_SEPARATOR = "+"


def resolved_separator(separator: str) -> str:
    """Blank falls back to "+" at build time — what is remembered is what was typed."""
    return separator if separator else JOINT_FALLBACK_SEPARATOR


def _double(text: str):
    if not text or text != text.strip():
        return None
    try:
        return float(text)
    except ValueError:
        return None


def tidy_grid(layout: Layout, include_unassigned: bool = True) -> Grid:
    """One row per well: [Plate] Well Row Column <factor…> [Note]. `Plate` only with >1
    plate; `Note` only when any well anywhere is noted."""
    multi = len(layout.plates) > 1
    any_notes = any(p.well_notes for p in layout.plates)
    header: list[str] = (["Plate"] if multi else []) + ["Well", "Row", "Column"] + [f.display_name for f in layout.factors]
    if any_notes:
        header.append("Note")
    grid: Grid = [header]
    for plate in layout.plates:
        fmt = plate.format
        for r in range(fmt.rows):
            for c in range(fmt.cols):
                well = fmt.index(r, c)
                values = []
                for f in layout.factors:
                    lid = plate.level_id(f.id, well)
                    lv = f.level(lid) if lid else None
                    values.append(lv.name if lv else "")
                if not include_unassigned and all(v == "" for v in values):
                    continue
                row: list[str] = [plate.name] if multi else []
                row += [WellNaming.well_label(r, c, layout.pad_well_labels), WellNaming.row_label(r), str(c + 1)]
                row += values
                if any_notes:
                    row.append(plate.note_for(well) or "")
                grid.append(row)
    return grid


def unique_sheet_names(names: list[str]) -> list[str]:
    """Excel sheet names: no []:*?/\\, non-empty, ≤31 chars, unique case-insensitively —
    de-duplicated with " (2)", " (3)"… trimming the base to fit (XLSX.uniquelyNamed)."""
    used: set[str] = set()
    out = []
    for raw in names:
        name = "".join(ch for ch in raw if ch not in "[]:*?/\\") or "Sheet"
        name = name[:31]
        candidate = name
        n = 2
        while candidate.lower() in used:
            suffix = f" ({n})"
            candidate = name[: 31 - len(suffix)] + suffix
            n += 1
        used.add(candidate.lower())
        out.append(candidate)
    return out


def needs_light_text(hex_text: str) -> bool:
    rgb = palette.parse_hex(hex_text)
    return rgb is not None and palette.luminance(rgb) <= 0.42


# ---- a tiny cell model so the sheet builders stay independent of openpyxl -------------


class Cell:
    __slots__ = ("value", "fill_hex", "bold", "centered")

    def __init__(self, value=None, fill_hex: str | None = None, bold: bool = False, centered: bool = False):
        self.value = value  # None (blank) | str | float/int
        self.fill_hex = fill_hex
        self.bold = bold
        self.centered = centered

    @classmethod
    def text(cls, s: str) -> "Cell":
        return cls(s)

    @classmethod
    def header(cls, s: str) -> "Cell":
        return cls(s, "#EDEDED", True, True)

    @classmethod
    def blank(cls) -> "Cell":
        return cls(None)


class Sheet:
    def __init__(self, name: str, rows: list, column_widths: list[float], freeze_rows: int = 0, freeze_cols: int = 0):
        self.name = name
        self.rows = rows  # list of list[Cell]; an empty list = a genuinely empty row
        self.column_widths = column_widths
        self.freeze_rows = freeze_rows
        self.freeze_cols = freeze_cols


def _map_rows(plate: Plate, factor: Factor) -> list:
    fmt = plate.format
    rows = [[Cell.header("")] + [Cell.header(str(c + 1)) for c in range(fmt.cols)]]
    for r in range(fmt.rows):
        row = [Cell.header(WellNaming.row_label(r))]
        for c in range(fmt.cols):
            lid = plate.level_id(factor.id, fmt.index(r, c))
            lv = factor.level(lid) if lid else None
            if lv is not None:
                value = lv.name
                if factor.kind is FactorKind.numeric:
                    num = _double(lv.name)
                    value = num if num is not None else lv.name
                row.append(Cell(value, lv.color_hex, False, True))
            else:
                row.append(Cell(None, None, False, True))
        rows.append(row)
    return rows


def _widths(fmt) -> list[float]:
    return [5.0] + [11.0] * fmt.cols


def _map_sheet(plate: Plate, factor: Factor) -> Sheet:
    return Sheet(f"{plate.name} · {factor.name}", _map_rows(plate, factor), _widths(plate.format), 1, 1)


def _combined_sheet(plate: Plate, factors) -> Sheet:
    rows: list = []
    for i, f in enumerate(factors):
        if i > 0:
            rows.append([])
        rows.append([Cell(f.display_name, None, True, False)])
        rows.extend(_map_rows(plate, f))
    return Sheet(plate.name, rows, _widths(plate.format), 0, 1)


def _joint_sheet(plate: Plate, factors, separator: str) -> Sheet:
    fmt = plate.format
    rows = [[Cell.header("")] + [Cell.header(str(c + 1)) for c in range(fmt.cols)]]
    for r in range(fmt.rows):
        row = [Cell.header(WellNaming.row_label(r))]
        for c in range(fmt.cols):
            well = fmt.index(r, c)
            parts = []
            for f in factors:
                lid = plate.level_id(f.id, well)
                lv = f.level(lid) if lid else None
                if lv is not None:
                    parts.append(lv.name)
            joined = separator.join(parts)
            row.append(Cell(joined if joined else None, None, False, True))
        rows.append(row)
    return Sheet(f"{plate.name} · Combined", rows, [5.0] + [24.0] * fmt.cols, 1, 1)


def _tidy_sheet(layout: Layout) -> Sheet:
    grid = tidy_grid(layout)
    numeric_columns = {i for i, f in enumerate(layout.factors) if f.kind is FactorKind.numeric}
    leading = (1 if len(layout.plates) > 1 else 0) + 3
    rows: list = []
    for i, row in enumerate(grid):
        if i == 0:
            rows.append([Cell.header(v) for v in row])
            continue
        cells = []
        for col, text in enumerate(row):
            num = _double(text) if (col >= leading and (col - leading) in numeric_columns) else None
            cells.append(Cell(num) if num is not None else Cell.text(text))
        rows.append(cells)
    widths = [float(max(9, min(24, len(h) + 4))) for h in grid[0]] if grid else []
    return Sheet("Wells", rows, widths, 1, 0)


def _legend_sheet(layout: Layout) -> Sheet:
    rows: list = [[Cell.header("Factor"), Cell.header("Level"), Cell.header("Colour"), Cell.header("Wells")]]
    for f in layout.factors:
        for lv in f.levels:
            count = sum(p.assigned_well_count(f.id, lv.id) for p in layout.plates)
            rows.append([Cell.text(f.display_name), Cell(lv.name, lv.color_hex, False, False), Cell.text(lv.color_hex), Cell(count)])
    noted = [p for p in layout.plates if p.note]
    if noted:
        rows.append([])
        rows.append([Cell("Plate notes", None, True, False)])
        for p in noted:
            rows.append([Cell.text(p.name), Cell.text(p.note)])
    return Sheet("Legend", rows, [22.0, 22.0, 12.0, 8.0], 1, 0)


def workbook_sheets(layout: Layout, sheet_layout: WorkbookLayout = WorkbookLayout.sheetPerFactor,
                    only_plate: str | None = None, joint_separator: str | None = None) -> list:
    """The sheets of an export, in order (`Exporter.workbook`). The scope is one filter
    applied first so the maps, Wells and Legend can never disagree; an id that matches
    nothing keeps the whole document."""
    if only_plate is not None:
        kept = tuple(p for p in layout.plates if p.id == only_plate)
        if kept:
            from dataclasses import replace

            layout = replace(layout, plates=kept)
    sheets: list = []
    if sheet_layout is WorkbookLayout.sheetPerFactor:
        for plate in layout.plates:
            for f in layout.factors:
                sheets.append(_map_sheet(plate, f))
    else:
        for plate in layout.plates:
            sheets.append(_combined_sheet(plate, layout.factors))
    if joint_separator is not None:
        for plate in layout.plates:
            sheets.append(_joint_sheet(plate, layout.factors, joint_separator))
    sheets.append(_tidy_sheet(layout))
    sheets.append(_legend_sheet(layout))
    names = unique_sheet_names([s.name for s in sheets])
    for s, n in zip(sheets, names):
        s.name = n
    return sheets


def build_workbook(layout: Layout, sheet_layout: WorkbookLayout = WorkbookLayout.sheetPerFactor,
                   only_plate: str | None = None, joint_separator: str | None = None):
    """An openpyxl Workbook with the Mac workbook's sheets, cell layout, fills, fonts,
    borders, widths and frozen panes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="FFD0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centre = Alignment(horizontal="center", vertical="center")
    fills: dict[str, PatternFill] = {}
    fonts: dict[tuple[bool, bool], Font] = {}

    def fill_for(hex_text: str) -> PatternFill:
        key = palette.normalized(hex_text)
        if key not in fills:
            rgb = key[1:] if key.startswith("#") else key
            fills[key] = PatternFill(patternType="solid", fgColor=f"FF{rgb}", bgColor=f"FF{rgb}")
        return fills[key]

    def font_for(bold: bool, light: bool) -> Font:
        key = (bold, light)
        if key not in fonts:
            fonts[key] = Font(name="Calibri", size=11, bold=bold, color="FFFFFFFF" if light else None)
        return fonts[key]

    for sheet in workbook_sheets(layout, sheet_layout, only_plate, joint_separator):
        ws = wb.create_sheet(title=sheet.name)
        for r, row in enumerate(sheet.rows, start=1):
            for c, cell in enumerate(row, start=1):
                if cell.value is None and cell.fill_hex is None and not cell.bold and not cell.centered:
                    continue
                target = ws.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    v = cell.value
                    target.value = int(v) if float(v).is_integer() else float(v)
                elif cell.value is not None:
                    target.value = cell.value
                    target.data_type = "s"  # never let a value starting with "=" become a formula
                light = cell.fill_hex is not None and needs_light_text(cell.fill_hex)
                if cell.bold or light:
                    target.font = font_for(cell.bold, light)
                if cell.fill_hex is not None:
                    target.fill = fill_for(cell.fill_hex)
                    target.border = border
                if cell.centered:
                    target.alignment = centre
        for i, width in enumerate(sheet.column_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        if sheet.freeze_rows or sheet.freeze_cols:
            ws.freeze_panes = ws.cell(row=sheet.freeze_rows + 1, column=sheet.freeze_cols + 1)
    return wb


def workbook_bytes(layout: Layout, sheet_layout: WorkbookLayout = WorkbookLayout.sheetPerFactor,
                   only_plate: str | None = None, joint_separator: str | None = None) -> bytes:
    import io

    buf = io.BytesIO()
    build_workbook(layout, sheet_layout, only_plate, joint_separator).save(buf)
    return buf.getvalue()


def tidy_csv(layout: Layout) -> str:
    """`Export Tidy CSV…` — the whole document, no scope choice."""
    return CSV.serialize(tidy_grid(layout))
