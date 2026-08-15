"""Exports — ports of CoreTests.TableIOTests (tidy grid) and WorkbookTests (sheet names,
arrangement, scope, joint map, cell fills/values), plus PNG/PDF/print smoke tests. The
workbook is verified by re-opening it with openpyxl, never by bytes."""
import io
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from playout.io.table_io import (
    CSV,
    WorkbookLayout,
    WorkbookScope,
    resolved_separator,
    tidy_csv,
    tidy_grid,
    unique_sheet_names,
    workbook_bytes,
    workbook_sheets,
)
from playout.model import palette
from playout.model.layout import Factor, Layout, Level, Plate
from playout.model.plate_format import WELL96, PlateFormat
from tests.ui_helpers import make_editor


def sample_layout():
    """WorkbookTests.sampleLayout: two factors incl. a numeric Dose (µM), two painted wells."""
    condition = Factor(name="Condition", levels=(Level("Untreated", palette.color_at(0)), Level("Treated", palette.color_at(1))))
    dose = Factor(name="Dose", kind="numeric", unit="µM", levels=(Level("10", "#5889BC"), Level("1", "#A0CBE8")))
    plate = Plate(name="Plate 1", format=WELL96)
    plate = plate.set_level_ids(condition.id, {0: condition.levels[0].id, 1: condition.levels[1].id})
    plate = plate.set_level_ids(dose.id, {0: dose.levels[0].id, 1: dose.levels[1].id})
    return Layout(factors=(condition, dose), plates=(plate,))


def open_wb(data: bytes):
    return load_workbook(io.BytesIO(data))


# ---------------------------------------------------------------- tidy grid / CSV


def test_tidy_grid_has_one_row_per_well_with_empty_strings_for_unassigned():
    grid = tidy_grid(sample_layout())
    assert grid[0] == ["Well", "Row", "Column", "Condition", "Dose (µM)"]
    assert len(grid) == 97
    assert grid[1] == ["A1", "A", "1", "Untreated", "10"]
    assert grid[3] == ["A3", "A", "3", "", ""]


def test_tidy_grid_names_plates_only_when_there_are_several_and_adds_note_column_only_when_noted():
    layout = sample_layout()
    two = replace(layout, plates=layout.plates + (Plate(name="Plate 2", format=PlateFormat(2, 3)),))
    grid = tidy_grid(two)
    assert grid[0][0] == "Plate" and len(grid) == 1 + 96 + 6
    assert grid[97] == ["Plate 2", "A1", "A", "1", "", ""]
    noted = replace(layout, plates=(layout.plates[0].set_note("bubble", 5),))
    grid = tidy_grid(noted)
    assert grid[0][-1] == "Note" and grid[6][-1] == "bubble" and grid[1][-1] == ""


def test_tidy_grid_respects_padded_labels_and_can_skip_unassigned():
    layout = replace(sample_layout(), pad_well_labels=True)
    grid = tidy_grid(layout)
    assert grid[1][0] == "A01"
    assert len(tidy_grid(layout, include_unassigned=False)) == 3


def test_tidy_csv_is_the_serialized_grid_with_a_trailing_newline():
    text = tidy_csv(sample_layout())
    assert text.startswith("Well,Row,Column,Condition,Dose (µM)\nA1,A,1,Untreated,10\n") and text.endswith("\n")
    assert CSV.parse(text) == tidy_grid(sample_layout())


# ---------------------------------------------------------------- workbook


def test_workbook_has_a_sheet_per_factor_plus_wells_and_legend():
    wb = open_wb(workbook_bytes(sample_layout()))
    assert wb.sheetnames == ["Plate 1 · Condition", "Plate 1 · Dose", "Wells", "Legend"]
    ws = wb["Plate 1 · Condition"]
    assert ws["B1"].value == 1 or ws["B1"].value == "1"
    assert ws["A2"].value == "A" and ws["B2"].value == "Untreated" and ws["C2"].value == "Treated"
    assert ws["B2"].fill.fgColor.rgb.upper() == "FF" + palette.color_at(0)[1:]
    assert ws["D2"].value is None and ws["D2"].fill.patternType is None
    assert ws.freeze_panes == "B2"
    assert ws["A1"].fill.fgColor.rgb.upper() == "FFEDEDED" and ws["A1"].font.bold
    dose = wb["Plate 1 · Dose"]
    assert dose["B2"].value == 10 and isinstance(dose["B2"].value, int)      # numeric factor → number
    wells = wb["Wells"]
    assert [c.value for c in wells[1]] == ["Well", "Row", "Column", "Condition", "Dose (µM)"]
    assert wells["E2"].value == 10 and wells["D2"].value == "Untreated"
    assert wells.freeze_panes == "A2"
    legend = wb["Legend"]
    assert [c.value for c in legend[1]] == ["Factor", "Level", "Colour", "Wells"]
    assert legend["A2"].value == "Condition" and legend["B2"].value == "Untreated" and legend["D2"].value == 1
    assert legend["B2"].fill.fgColor.rgb.upper() == "FF" + palette.color_at(0)[1:]
    assert legend["A4"].value == "Dose (µM)" and legend["C4"].value == "#5889BC"


def test_combined_layout_uses_one_sheet_per_plate_headed_by_factor_names():
    wb = open_wb(workbook_bytes(sample_layout(), WorkbookLayout.allFactorsOneSheet))
    assert wb.sheetnames == ["Plate 1", "Wells", "Legend"]
    ws = wb["Plate 1"]
    assert ws["A1"].value == "Condition" and ws["A1"].font.bold
    assert ws["A2"].value is None and ws["B2"].value in (1, "1")     # header row of the first map
    # block pitch: 1 title + 1 header + 8 rows + 1 gap → the second title at row 12
    assert ws["A12"].value == "Dose (µM)" and ws["A11"].value is None
    assert ws.freeze_panes == "B1"


def test_joint_map_is_off_by_default_and_joins_every_factor_in_one_cell():
    layout = sample_layout()
    assert "Plate 1 · Combined" not in open_wb(workbook_bytes(layout)).sheetnames
    wb = open_wb(workbook_bytes(layout, joint_separator=" / "))
    assert wb.sheetnames == ["Plate 1 · Condition", "Plate 1 · Dose", "Plate 1 · Combined", "Wells", "Legend"]
    ws = wb["Plate 1 · Combined"]
    assert ws["B2"].value == "Untreated / 10" and ws["C2"].value == "Treated / 1" and ws["D2"].value is None
    assert ws["B2"].fill.patternType is None
    assert resolved_separator("") == "+" and resolved_separator("|") == "|"


def test_workbook_can_cover_just_one_plate_and_an_unmatched_id_keeps_every_plate():
    layout = sample_layout()
    p2 = Plate(name="Plate 2", format=PlateFormat(2, 3))
    two = replace(layout, plates=layout.plates + (p2,))
    only = open_wb(workbook_bytes(two, only_plate=p2.id))
    assert only.sheetnames == ["Plate 2 · Condition", "Plate 2 · Dose", "Wells", "Legend"]
    assert only["Wells"].max_row == 1 + 6 and only["Wells"]["A1"].value == "Well"   # one plate left → no Plate column
    every = open_wb(workbook_bytes(two, only_plate="AAAAAAAA-BBBB-4CCC-8DDD-EEEEEEEEEEEE"))
    assert len(every.sheetnames) == 6


def test_legend_carries_plate_notes_and_light_text_on_dark_fills():
    layout = sample_layout()
    dark = Factor(name="Dark", levels=(Level("ink", "#102A44"),))
    plate = layout.plates[0].with_note("edge effects").set_level_id(dark.levels[0].id, dark.id, 0)
    layout = replace(layout, factors=layout.factors + (dark,), plates=(plate,))
    wb = open_wb(workbook_bytes(layout))
    legend = wb["Legend"]
    rows = [[c.value for c in r] for r in legend.iter_rows()]
    assert ["Plate notes", None, None, None] in rows or ["Plate notes"] == [v for v in rows[-2] if v]
    assert rows[-1][:2] == ["Plate 1", "edge effects"]
    ink = wb["Plate 1 · Dark"]["B2"]
    assert ink.font.color is not None and ink.font.color.rgb.upper() == "FFFFFFFF"
    # the Mac rule is luminance ≤ 0.42 → white, so a pale yellow keeps the default ink
    pale = Factor(name="Pale", levels=(Level("y", "#EDC948"),))
    layout2 = replace(layout, factors=(pale,), plates=(layout.plates[0].set_level_id(pale.levels[0].id, pale.id, 0),))
    cell = open_wb(workbook_bytes(layout2))["Plate 1 · Pale"]["B2"]
    assert cell.font.color is None or str(cell.font.color.rgb).upper() != "FFFFFFFF"


def test_sheet_names_are_sanitised_truncated_and_deduplicated():
    assert unique_sheet_names(["a/b:c", "a[b]c"]) == ["abc", "abc (2)"]
    assert unique_sheet_names(["x" * 40, "x" * 40, "X" * 40]) == ["x" * 31, "x" * 27 + " (2)", "X" * 27 + " (3)"]
    assert unique_sheet_names([""]) == ["Sheet"]
    two = sample_layout()
    two = replace(two, plates=(two.plates[0], replace(two.plates[0], id="AAAAAAAA-BBBB-4CCC-8DDD-EEEEEEEEEEEE")))
    names = [s.name for s in workbook_sheets(two)]
    assert names[:4] == ["Plate 1 · Condition", "Plate 1 · Dose", "Plate 1 · Condition (2)", "Plate 1 · Dose (2)"]


def test_a_value_starting_with_equals_is_stored_as_text():
    f = Factor(name="F", levels=(Level("=SUM(1)", "#5889BC"),))
    layout = Layout(factors=(f,), plates=(Plate(name="P", format=PlateFormat(1, 1)).set_level_id(f.levels[0].id, f.id, 0),))
    ws = open_wb(workbook_bytes(layout))["P · F"]
    assert ws["B2"].value == "=SUM(1)" and ws["B2"].data_type == "s"


def test_workbook_from_the_mac_fixture_round_trips_every_well(fixtures_dir):
    from playout.model.layout import loads

    layout = loads((fixtures_dir / "allfactors.plate").read_text(encoding="utf-8"))
    wb = open_wb(workbook_bytes(layout, joint_separator="+"))
    wells = wb["Wells"]
    total = sum(p.format.well_count for p in layout.plates)
    assert wells.max_row == 1 + total and wells["A1"].value == "Plate"


# ---------------------------------------------------------------- image / pdf / print


@pytest.fixture(autouse=True)
def _gui(qapp):
    pass


def test_png_export_is_a_png_at_the_requested_size(tmp_path):
    from PySide6.QtGui import QImage

    from playout.io.plate_image import render_png

    ed = make_editor(tmp_path)
    data = render_png(ed, 940, 560, 2.0)
    assert data is not None and data[:8] == b"\x89PNG\r\n\x1a\n"
    img = QImage.fromData(data, "PNG")
    assert img.width() == 1880 and img.height() == 1120
    assert render_png(ed, 2, 2) is None


def test_pdf_export_writes_a_pdf_and_print_paints_to_a_pdf_printer(tmp_path):
    from PySide6.QtPrintSupport import QPrinter

    from playout.io.plate_image import paint_for_print, write_pdf

    ed = make_editor(tmp_path)
    ed.select_all_wells()
    ed.paint_selection()
    pdf = tmp_path / "plate.pdf"
    assert write_pdf(ed, str(pdf), 940, 560)
    assert pdf.read_bytes()[:5] == b"%PDF-"
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    out = tmp_path / "print.pdf"
    printer.setOutputFileName(str(out))
    assert paint_for_print(ed, printer, 940, 560)
    assert out.read_bytes()[:5] == b"%PDF-"


def test_workbook_options_dialog_remembers_choices(tmp_path, qtbot):
    from playout.ui.sheets.workbook_options import WorkbookOptionsDialog

    ed = make_editor(tmp_path)
    dlg = WorkbookOptionsDialog(ed.preferences, plate_count=1, active_plate_name="Plate 1")
    qtbot.addWidget(dlg)
    assert not dlg.scope.isVisibleTo(dlg) and dlg.selected_scope is WorkbookScope.allPlates
    dlg.arrangement.setCurrentIndex(1)
    dlg.joint.setChecked(True)
    dlg.separator.setText("|")
    dlg.remember()
    assert ed.preferences.workbook_sheet_layout == "allFactorsOneSheet"
    assert ed.preferences.workbook_joint_map_enabled and ed.preferences.workbook_joint_map_separator == "|"
    dlg2 = WorkbookOptionsDialog(ed.preferences, plate_count=2, active_plate_name="A very long plate name indeed, truly")
    qtbot.addWidget(dlg2)
    assert dlg2.selected_layout is WorkbookLayout.allFactorsOneSheet and dlg2.joint.isChecked()
    assert dlg2.scope.isVisibleTo(dlg2) and dlg2.scope.itemText(1) == 'Just "A very long plate name in…"'
